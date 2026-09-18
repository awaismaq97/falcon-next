"""
text_extract.py — pull plain text out of an uploaded or downloaded document.

This logic used to live as private functions inside ``app/routers/documents.py``,
which was fine while a browser upload was the only way a document entered the
system. It no longer is: ``falcon.google_drive`` downloads files from Drive and
needs exactly the same extraction, and a domain module importing a FastAPI
router to get at its underscore-prefixed helpers would be backwards.

So the extractors live here, in the domain layer, and the router is a thin caller
that maps :class:`UnsupportedDocument` onto an HTTP status. Nothing about the
extraction itself changed — the same libraries, the same per-format handling, the
same last-resort text decode.

``UnsupportedDocument`` deliberately derives from ``Exception`` rather than
``ValueError``: ``app.main`` installs a global ValueError handler that answers
400, and "this file type is not supported" is a 415. Inheriting from ValueError
would silently route it to the wrong status from anywhere that did not catch it.
"""
from __future__ import annotations

import io
import logging

logger = logging.getLogger("falcon.text_extract")

# Text-like formats decoded directly rather than parsed.
TEXT_EXTS = frozenset({
    "txt", "md", "markdown", "csv", "tsv", "json", "xml", "html", "htm",
    "yaml", "yml", "log", "ini", "cfg", "rtf", "tex",
})

# Formats with a real parser behind them.
BINARY_EXTS = frozenset({"pdf", "docx", "xlsx", "xlsm", "pptx"})

# Everything this module can read, for callers that want to filter before
# spending a download.
SUPPORTED_EXTS = TEXT_EXTS | BINARY_EXTS

# The pre-2007 binary Office formats. Named separately so the error can say what
# to do about it instead of "unsupported".
LEGACY_EXTS = frozenset({"doc", "xls", "ppt"})


class UnsupportedDocument(Exception):
    """The file is not a format this module can read."""


def ext_of(name: str) -> str:
    """The lowercased extension of a filename, or "" when it has none."""
    return name.rsplit(".", 1)[-1].lower() if "." in (name or "") else ""


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")  # try empty password
        except Exception:  # noqa: BLE001
            pass
    parts = []
    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            txt = ""
        if txt.strip():
            parts.append(txt)
    return "\n\n".join(parts)


def _extract_docx(data: bytes) -> str:
    import docx  # python-docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                parts.append("\t".join(cells))
    try:
        wb.close()
    except Exception:  # noqa: BLE001
        pass
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"# Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        parts.append(line)
    return "\n".join(parts)


def decode_text(data: bytes) -> str:
    """Decode bytes that are already text, trying the encodings that show up."""
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract(name: str, data: bytes, ext: str = "") -> str:
    """Plain text from one document's bytes.

    ``ext`` may be passed when the caller knows the format from something other
    than the filename — a Drive MIME type, say — and is otherwise derived from
    ``name``.

    Raises:
        UnsupportedDocument: the format has no extractor, or is a legacy binary
            Office format that the user should re-save.
    """
    ext = (ext or ext_of(name)).lower().lstrip(".")

    if ext == "pdf":
        return _extract_pdf(data)
    if ext == "docx":
        return _extract_docx(data)
    if ext in ("xlsx", "xlsm"):
        return _extract_xlsx(data)
    if ext == "pptx":
        return _extract_pptx(data)
    if ext in TEXT_EXTS:
        return decode_text(data)
    if ext in LEGACY_EXTS:
        raise UnsupportedDocument(
            f"Legacy .{ext} files aren't supported — re-save as .{ext}x and try again."
        )

    # Unknown extension: attempt a text decode, and refuse if it looks binary.
    text = decode_text(data)
    if "\x00" in text[:1000]:
        raise UnsupportedDocument(f"Unsupported file type: .{ext or '?'}")
    return text
